"""Tests for FinalBidWordGenerator."""
import pytest
import tempfile
import os
from unittest.mock import MagicMock, patch


class TestFinalBidWordGenerator:
    """TDD tests for FinalBidWordGenerator."""

    def test_generate_creates_docx_file(self):
        """Test that generate_final_document creates a .docx file."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.docx")

            mock_db = MagicMock()
            mock_project = MagicMock()
            mock_project.id = 1
            mock_project.project_name = "测试项目"
            mock_project.budget_amount = 1500000.0
            mock_project.owner_unit = "测试单位"
            mock_db.get.return_value = mock_project

            tech_sections = {
                "sections": [
                    {"section_title": "项目概况", "content": "测试内容", "score_weight": "10"}
                ]
            }
            pricing_data = {
                "cost_base": 100000.0,
                "boss_final_price": 120000.0
            }

            from app.core.week5_formal_review.word_generator import FinalBidWordGenerator
            result = FinalBidWordGenerator.generate_final_document(
                project_id=1,
                tech_sections=tech_sections,
                pricing_data=pricing_data,
                output_path=output_path,
                db=mock_db
            )

            assert os.path.exists(result['file_path'])
            assert result['file_path'].endswith('.docx')
            assert result['status'] == 'completed'

    def test_generated_docx_contains_project_name(self):
        """Test that the generated doc contains the project name."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.docx")

            mock_db = MagicMock()
            mock_project = MagicMock()
            mock_project.id = 1
            mock_project.project_name = "测试项目XYZ"
            mock_project.budget_amount = 1500000.0
            mock_project.owner_unit = "测试单位"
            mock_db.get.return_value = mock_project

            tech_sections = {"sections": []}
            pricing_data = {"cost_base": 100000.0, "boss_final_price": 120000.0}

            from app.core.week5_formal_review.word_generator import FinalBidWordGenerator
            FinalBidWordGenerator.generate_final_document(
                project_id=1,
                tech_sections=tech_sections,
                pricing_data=pricing_data,
                output_path=output_path,
                db=mock_db
            )

            from docx import Document
            doc = Document(output_path)
            full_text = '\n'.join([p.text for p in doc.paragraphs])
            assert "测试项目XYZ" in full_text

    def test_pricing_table_in_document(self):
        """Test that pricing table is generated with correct values."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.docx")

            mock_db = MagicMock()
            mock_project = MagicMock()
            mock_project.id = 1
            mock_project.project_name = "测试项目"
            mock_project.budget_amount = 1500000.0
            mock_project.owner_unit = "测试单位"
            mock_db.get.return_value = mock_project

            tech_sections = {"sections": []}
            pricing_data = {
                "cost_base": 100000.0,
                "boss_final_price": 120000.0,
                "system_suggested_low": 110000.0,
                "system_suggested_high": 130000.0,
                "budget_limit": 140000.0
            }

            from app.core.week5_formal_review.word_generator import FinalBidWordGenerator
            FinalBidWordGenerator.generate_final_document(
                project_id=1,
                tech_sections=tech_sections,
                pricing_data=pricing_data,
                output_path=output_path,
                db=mock_db
            )

            from docx import Document
            doc = Document(output_path)
            tables = doc.tables
            assert len(tables) >= 1
            # Check table contains cost_base and boss_final_price
            table_text = ''
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        table_text += cell.text
            assert "成本基准" in table_text
            assert "120000" in table_text

    def test_tech_sections_headings_in_document(self):
        """Test that tech sections are added as Heading 2."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.docx")

            mock_db = MagicMock()
            mock_project = MagicMock()
            mock_project.id = 1
            mock_project.project_name = "测试项目"
            mock_project.budget_amount = 1500000.0
            mock_project.owner_unit = "测试单位"
            mock_db.get.return_value = mock_project

            tech_sections = {
                "sections": [
                    {"section_title": "第一章节", "content": "内容一", "score_weight": "10"},
                    {"section_title": "第二章节", "content": "内容二", "score_weight": "15"}
                ]
            }
            pricing_data = {"cost_base": 100000.0, "boss_final_price": 120000.0}

            from app.core.week5_formal_review.word_generator import FinalBidWordGenerator
            FinalBidWordGenerator.generate_final_document(
                project_id=1,
                tech_sections=tech_sections,
                pricing_data=pricing_data,
                output_path=output_path,
                db=mock_db
            )

            from docx import Document
            doc = Document(output_path)
            # Check for Heading 2 styles
            heading2_paras = [p for p in doc.paragraphs if p.style.name == 'Heading 2']
            assert len(heading2_paras) >= 2

    def test_packaging_guide_returned(self):
        """Test that packaging guide is returned in the result."""
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = os.path.join(tmpdir, "test_output.docx")

            mock_db = MagicMock()
            mock_project = MagicMock()
            mock_project.id = 1
            mock_project.project_name = "测试项目"
            mock_project.budget_amount = 1500000.0
            mock_project.owner_unit = "测试单位"
            mock_db.get.return_value = mock_project

            tech_sections = {"sections": []}
            pricing_data = {"cost_base": 100000.0, "boss_final_price": 120000.0}

            from app.core.week5_formal_review.word_generator import FinalBidWordGenerator
            result = FinalBidWordGenerator.generate_final_document(
                project_id=1,
                tech_sections=tech_sections,
                pricing_data=pricing_data,
                output_path=output_path,
                db=mock_db
            )

            assert "packaging_guide" in result
            assert "documents_checklist" in result["packaging_guide"]
            assert isinstance(result["packaging_guide"]["documents_checklist"], list)
