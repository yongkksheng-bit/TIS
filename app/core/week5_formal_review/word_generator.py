"""FinalBidWordGenerator — assembles final bid document using python-docx."""
import os
import uuid
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from sqlalchemy.orm import Session


class FinalBidWordGenerator:
    """Generates final bid Word document with tech sections and pricing table."""

    PACKAGING_GUIDE = {
        "seal_bags": [
            {"type": "正本", "copies": 1, "label": "技术标正本"},
            {"type": "副本", "copies": 1, "label": "技术标副本"}
        ],
        "documents_checklist": [
            "营业执照复印件",
            "资质证书复印件",
            "授权书原件",
            "项目经理证书复印件"
        ],
        "special_notes": "所有副本需加盖骑缝章"
    }

    @staticmethod
    def generate_final_document(
        project_id: int,
        tech_sections: dict,
        pricing_data: dict,
        output_path: str,
        db: Session = None
    ) -> dict:
        """
        Generate the final bid Word document.

        Args:
            project_id: Project ID to look up project details
            tech_sections: Dict with 'sections' key containing list of
                          {section_title, content, score_weight}
            pricing_data: Dict with cost_base, boss_final_price, and optional
                         system_suggested_low, system_suggested_high, budget_limit, decision_reason
            output_path: Full path for output .docx file

        Returns:
            dict with file_path, file_size, status, packaging_guide
        """
        # Ensure output directory exists
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

        # Look up project name from database
        project_name = f"项目 #{project_id}"
        if db is not None:
            from app.models.project import Project
            project = db.get(Project, project_id)
            if project:
                project_name = project.project_name

        doc = Document()
        cover_heading = doc.add_heading(project_name, level=0)
        cover_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 2. Project info paragraph
        budget = pricing_data.get("budget_limit", tech_sections.get("budget", "未设置"))
        owner_unit = "招标单位"
        date_str = datetime.now().strftime("%Y年%m月%d日")

        info_para = doc.add_paragraph()
        info_para.add_run(f"预算金额: {budget}\n").bold = True
        info_para.add_run(f"业主单位: {owner_unit}\n").bold = True
        info_para.add_run(f"日期: {date_str}").bold = True
        info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 3. Page break
        doc.add_page_break()

        # 4. Section heading "技术标" (Heading 1)
        doc.add_heading("技术标", level=1)

        # 5. For each tech section: Heading 2 + paragraph with content
        sections = tech_sections.get("sections", [])
        for section in sections:
            section_title = section.get("section_title", "未命名章节")
            content = section.get("content", "")
            doc.add_heading(section_title, level=2)
            doc.add_paragraph(content)

        # 6. Page break
        doc.add_page_break()

        # 7. Section heading "商务标报价" (Heading 1)
        doc.add_heading("商务标报价", level=1)

        # 8. Pricing table (2-column, 5 rows)
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'

        # Row 0: 成本基准
        table.cell(0, 0).text = "成本基准"
        table.cell(0, 1).text = str(pricing_data.get("cost_base", ""))

        # Row 1: 建议低价
        table.cell(1, 0).text = "建议低价"
        table.cell(1, 1).text = str(pricing_data.get("system_suggested_low", ""))

        # Row 2: 建议高价
        table.cell(2, 0).text = "建议高价"
        table.cell(2, 1).text = str(pricing_data.get("system_suggested_high", ""))

        # Row 3: 最终报价 (BOLD)
        table.cell(3, 0).text = "最终报价"
        cell = table.cell(3, 1)
        run = cell.paragraphs[0].add_run(str(pricing_data.get("boss_final_price", "")))
        run.bold = True

        # Row 4: 预算限价
        table.cell(4, 0).text = "预算限价"
        table.cell(4, 1).text = str(pricing_data.get("budget_limit", ""))

        # 9. If decision_reason exists, add paragraph
        decision_reason = pricing_data.get("decision_reason")
        if decision_reason:
            doc.add_paragraph()
            reason_para = doc.add_paragraph()
            reason_para.add_run("报价决策说明: ").bold = True
            reason_para.add_run(decision_reason)

        # 10. Page break
        doc.add_page_break()

        # 11. Section heading "封装指南" (Heading 1)
        doc.add_heading("封装指南", level=1)

        # 12. Checklist paragraphs with □ characters
        checklist_items = [
            "技术标正本 1 份",
            "技术标副本 1 份",
            "商务标正本 1 份",
            "商务标副本 1 份",
            "营业执照复印件",
            "资质证书复印件",
            "授权书原件",
            "项目经理证书复印件",
            "所有副本需加盖骑缝章"
        ]

        for item in checklist_items:
            doc.add_paragraph(f"□ {item}")

        # Save the document
        doc.save(output_path)

        # Get file size
        file_size = os.path.getsize(output_path)

        return {
            "file_path": output_path,
            "file_size": file_size,
            "status": "completed",
            "packaging_guide": FinalBidWordGenerator.PACKAGING_GUIDE
        }
